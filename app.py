import re
from html import escape
from io import BytesIO
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

# ============================================================
# ONLINE PROIZVODNJA DASHBOARD
# Verzija za Streamlit Cloud / online korišćenje preko upload-a Excel fajla
# ============================================================

st.set_page_config(page_title="Dnevni pregled proizvodnje", layout="wide")

# ============================================================
# MAPIRANJA IZ POSTOJEĆEG PROGRAMA
# ============================================================

PRVI_RED_PODATAKA = 8
POSLEDNJI_RED_PODATAKA = 40
RED_SMENA = 6
RED_NAZIVI = 7
KOLONA_DATUM = 1

DOZVOLJENE_SMENE = ["I shift", "II shift", "III shift"]
SMENA_LABEL = {
    "I shift": "1. smena",
    "II shift": "2. smena",
    "III shift": "3. smena",
}

PRESKOCI_TABOVE = [
    "AIDA 2 APP350",
    "Realizacija CW 21",
    "Realizacija CW 20",
    "Realizacija CW 22",
    "REALIZACIJA MESECNA",
    "Plan rada linija",
    "Realizacija smena",
    "Analiza skarta cw 18",
    "Analiza skarta cw 19",
    "Analiza skarta cw20",
    "Analiza skarta cw13",
    "Stops analysis CW7",
    "Sheet1",
    "Annealing",
]

MAPA_PROJEKATA = {
    "AIDA L1 APP350": "APP350",
    "AIDA 1 APP350": "APP350",
    "AIDA 3 APP350": "APP350",
    "AIDA 8 APP350": "APP350",
    "AIDA 9 APP350": "APP350",
    "AIDA 7 APP350": "APP350",
    "AIDA L7 APP350": "APP350",
    "DMC L2 APP350": "APP350",
    "DMC L3 APP350": "APP350",
    "DMC L4 APP350": "APP350",
    "Rotor L1 APP350": "APP350",
    "Rotor L3 APP350": "APP350",
    "Rotor L8 APP350": "APP350",

    "AIDA L2": "APP550",
    "AIDA 2": "APP550",
    "AIDA L4": "APP550",
    "AIDA 4": "APP550",
    "AIDA L5": "APP550",
    "AIDA 5": "APP550",
    "AIDA L6": "APP550",
    "AIDA 6": "APP550",
    "AIDA L7": "APP550",
    "AIDA 7": "APP550",
    "Stator L1": "APP550",
    "Stator L2": "APP550",
    "Stator L3": "APP550",
    "Stator L4": "APP550",
    "Stator L5": "APP550",
    "DMC L1 APP550": "APP550",
    "DMC L1 GP12": "APP550",
    "DMC L2 APP550": "APP550",
    "DMC L2 GP12": "APP550",
    "DMC L3 APP550": "APP550",
    "DMC L3 GP12": "APP550",
    "DMC L4 APP550": "APP550",
    "DMC L4 GP12": "APP550",
    "DMC L5 APP550": "APP550",
    "Rotor L2 APP550": "APP550",
    "Rotor L4 APP550": "APP550",
    "Rotor L5 APP550": "APP550",
    "Rotor L6 APP550": "APP550",

    "DMC EMR4 ROTOR": "VITESKO EMR4",
    "DMC EMR4 STATOR": "VITESKO EMR4",
    "HEATING VITESCO ROTOR": "VITESKO EMR4",
    "HEATING VITESCO STATOR": "VITESKO EMR4",
    "AIDA VITESKO": "VITESKO EMR4",

    "AIDA LK-4": "LK-4",
    "AIDA LK4": "LK-4",
    "AIDA LK 4": "LK-4",
    "AIDA LK-3": "LK-4",
    "DMC L1 LK-4": "LK-4",
    "DMC L1 LK4": "LK-4",
    "DMC L1 LK 4": "LK-4",
    "Rotor L10 LK-4": "LK-4",
    "Rotor L10 LK4": "LK-4",
    "Rotor L10 LK 4": "LK-4",
}

MAPA_PROCESA = {
    "AIDA L1 APP350": "STAMPING",
    "AIDA 1 APP350": "STAMPING",
    "AIDA L2": "STAMPING",
    "AIDA 2": "STAMPING",
    "AIDA 3 APP350": "STAMPING",
    "AIDA L4": "STAMPING",
    "AIDA 4": "STAMPING",
    "AIDA L5": "STAMPING",
    "AIDA 5": "STAMPING",
    "AIDA L6": "STAMPING",
    "AIDA 6": "STAMPING",
    "AIDA L7": "STAMPING",
    "AIDA 7": "STAMPING",
    "AIDA 8 APP350": "STAMPING",
    "AIDA 9 APP350": "STAMPING",
    "AIDA 7 APP350": "STAMPING",
    "AIDA L7 APP350": "STAMPING",
    "AIDA LK-4": "STAMPING",
    "AIDA LK4": "STAMPING",
    "AIDA LK 4": "STAMPING",
    "AIDA LK-3": "STAMPING",
    "AIDA VITESKO": "STAMPING",

    "Stator L1": "WELDING",
    "Stator L2": "WELDING",
    "Stator L3": "WELDING",
    "Stator L4": "WELDING",
    "Stator L5": "WELDING",
    "HEATING VITESCO STATOR": "WELDING",
    "HEATING VITESCO ROTOR": "WELDING",

    "DMC L1 APP550": "DMC",
    "DMC L2 APP550": "DMC",
    "DMC L3 APP550": "DMC",
    "DMC L4 APP550": "DMC",
    "DMC L5 APP550": "DMC",
    "DMC L2 APP350": "DMC",
    "DMC L3 APP350": "DMC",
    "DMC L4 APP350": "DMC",
    "DMC L1 LK-4": "DMC",
    "DMC L1 LK4": "DMC",
    "DMC L1 LK 4": "DMC",
    "DMC EMR4 STATOR": "DMC",
    "DMC EMR4 ROTOR": "DMC",

    "DMC L1 GP12": "GP12",
    "DMC L2 GP12": "GP12",
    "DMC L3 GP12": "GP12",
    "DMC L4 GP12": "GP12",

    "Rotor L1 APP350": "ROTOR",
    "Rotor L2 APP550": "ROTOR",
    "Rotor L3 APP350": "ROTOR",
    "Rotor L4 APP550": "ROTOR",
    "Rotor L5 APP550": "ROTOR",
    "Rotor L6 APP550": "ROTOR",
    "Rotor L8 APP350": "ROTOR",
    "Rotor L10 LK-4": "ROTOR",
    "Rotor L10 LK4": "ROTOR",
    "Rotor L10 LK 4": "ROTOR",
}

# Imena/autori notes-a koja ne smeju da postanu razlog zastoja ili NOK-a.
RECI_KOJE_SE_BRISU_IZ_RAZLOGA = [
    "bojan", "dimitrije", "smiljanić", "smiljanic", "miler",
    "jelena", "stefanović", "stefanovic",
    "valentina", "savatić", "savatic",
    "atila", "čolak", "colak",
]

PUNA_IMENA_ZA_BRISANJE = [
    "Bojan Smiljanić", "Bojan Smiljanic",
    "Dimitrije Miler",
    "Jelena Stefanović", "Jelena Stefanovic",
    "Valentina Savatić", "Valentina Savatic",
    "Atila Čolak", "Atila Colak",
]

# ============================================================
# CSS
# ============================================================

st.markdown(
    """
    <style>
    .stApp {
        background: linear-gradient(135deg, #0f172a 0%, #111827 45%, #020617 100%);
        color: #e5e7eb;
    }
    section[data-testid="stSidebar"] {
        background-color: #020617;
        border-right: 1px solid #1f2937;
        height: 100vh !important;
        max-height: 100vh !important;
        overflow-y: auto !important;
        overflow-x: hidden !important;
    }
    section[data-testid="stSidebar"] * { color: #e5e7eb !important; }
    section[data-testid="stSidebar"] input,
    section[data-testid="stSidebar"] div[data-baseweb="select"] input,
    section[data-testid="stSidebar"] div[data-baseweb="select"] span,
    section[data-testid="stSidebar"] div[data-baseweb="select"] div,
    section[data-testid="stSidebar"] [data-testid="stDateInput"] input {
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="tag"] span {
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
    }
    .block-container { padding-top: 1.2rem; max-width: 100%; }
    h1, h2, h3, h4 { color: #f8fafc !important; }
    .machine-card {
        background: rgba(15, 23, 42, 0.86);
        border: 1px solid rgba(148, 163, 184, 0.35);
        border-radius: 18px;
        padding: 18px;
        margin: 16px 0 24px 0;
        box-shadow: 0 12px 35px rgba(0,0,0,0.25);
    }
    .machine-title {
        font-size: 26px;
        font-weight: 900;
        color: #ffffff;
        margin-bottom: 4px;
    }
    .machine-subtitle {
        color: #cbd5e1;
        font-size: 14px;
        margin-bottom: 14px;
    }
    .mini-title {
        color: #f8fafc;
        font-size: 18px;
        font-weight: 800;
        margin-top: 12px;
        margin-bottom: 8px;
    }
    [data-testid="stMetric"] {
        background: rgba(2, 6, 23, 0.78);
        border: 1px solid rgba(51, 65, 85, 0.9);
        border-radius: 14px;
        padding: 12px;
    }
    [data-testid="stMetric"] * { color: #f8fafc !important; }
    [data-testid="stDataFrame"] {
        background: rgba(15, 23, 42, 0.88) !important;
    }
    .neo-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
        gap: 12px;
        margin: 12px 0 18px 0;
    }
    .neo-card {
        background: linear-gradient(145deg, rgba(15,23,42,0.94), rgba(2,6,23,0.96));
        border: 1px solid rgba(56,189,248,0.24);
        border-radius: 18px;
        padding: 14px 16px;
        box-shadow: 0 16px 35px rgba(0,0,0,0.30), inset 0 1px 0 rgba(255,255,255,0.06);
        position: relative;
        overflow: hidden;
    }
    .neo-card::before {
        content: "";
        position: absolute;
        top: -40px;
        right: -40px;
        width: 90px;
        height: 90px;
        background: radial-gradient(circle, rgba(34,211,238,0.22), transparent 70%);
    }
    .neo-card-total {
        border-color: rgba(34,197,94,0.45);
        background: linear-gradient(145deg, rgba(6,78,59,0.42), rgba(2,6,23,0.96));
    }
    .neo-label {
        color: #94a3b8;
        font-size: 12px;
        text-transform: uppercase;
        letter-spacing: .08em;
        font-weight: 800;
        margin-bottom: 6px;
    }
    .neo-value {
        color: #f8fafc;
        font-size: 24px;
        font-weight: 900;
        line-height: 1.05;
    }
    .neo-sub {
        color: #cbd5e1;
        font-size: 13px;
        margin-top: 7px;
    }
    .neo-pill {
        display: inline-block;
        background: rgba(14,165,233,0.14);
        border: 1px solid rgba(125,211,252,0.30);
        color: #e0f2fe;
        padding: 5px 9px;
        border-radius: 999px;
        font-size: 12px;
        font-weight: 800;
        margin: 3px 5px 3px 0;
    }
    .stop-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
        gap: 10px;
        margin-top: 8px;
    }
    .stop-card {
        background: rgba(2,6,23,0.78);
        border: 1px solid rgba(148,163,184,0.22);
        border-left: 4px solid rgba(56,189,248,0.85);
        border-radius: 14px;
        padding: 12px 14px;
    }
    .stop-reason {
        color: #ffffff;
        font-size: 16px;
        font-weight: 850;
        margin-bottom: 7px;
    }
    .stop-meta {
        color: #cbd5e1;
        font-size: 13px;
    }
    .summary-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
        gap: 12px;
        margin: 10px 0 18px 0;
    }
    .summary-card {
        background: linear-gradient(145deg, rgba(15,23,42,0.93), rgba(2,6,23,0.97));
        border: 1px solid rgba(148,163,184,0.25);
        border-radius: 18px;
        padding: 14px;
        box-shadow: 0 12px 28px rgba(0,0,0,0.24);
    }
    .summary-title {
        color: #f8fafc;
        font-size: 18px;
        font-weight: 900;
        margin-bottom: 6px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ============================================================
# POMOĆNE FUNKCIJE
# ============================================================

def ocisti_tekst(vrednost):
    if vrednost is None:
        return None
    tekst = str(vrednost).replace("\r", "\n").replace("\n", " ").strip()
    return tekst or None


def norm(vrednost):
    tekst = ocisti_tekst(vrednost)
    if tekst is None:
        return ""
    return " ".join(tekst.lower().split())


def broj(vrednost):
    if vrednost is None:
        return 0.0
    if isinstance(vrednost, str):
        vrednost = vrednost.replace(".", "").replace(",", ".") if re.search(r"\d+,\d+", vrednost) else vrednost
    try:
        return float(vrednost)
    except Exception:
        return 0.0


def format_broj(x):
    if x is None or pd.isna(x):
        return "-"
    return f"{float(x):,.0f}".replace(",", ".")


def format_proc(x):
    if x is None or pd.isna(x):
        return "-"
    return f"{float(x):.1f}%"


def procenat(deo, ukupno):
    if ukupno is None or ukupno == 0:
        return None
    return deo / ukupno * 100


def kljuc_masine(naziv):
    """Normalizuje naziv taba/mašine da LK4, LK-4 i LK 4 budu isto."""
    t = str(naziv or "").upper().strip()
    t = t.replace("Č", "C").replace("Ć", "C").replace("Š", "S").replace("Ž", "Z").replace("Đ", "DJ")
    t = re.sub(r"\s+", " ", t)
    t = re.sub(r"LK\s*-?\s*([34])", r"LK\1", t)
    t = re.sub(r"L\s*-?\s*(\d+)", r"L\1", t)
    t = t.replace("-", "")
    t = re.sub(r"\s+", " ", t).strip()
    return t


def pronadji_u_mapi(naziv, mapa, default="NEMAPIRANO"):
    if naziv is None:
        return default

    naziv_cist = str(naziv).strip()
    if naziv_cist in mapa:
        return mapa[naziv_cist]

    trazeni = kljuc_masine(naziv_cist)
    for kljuc, vrednost in mapa.items():
        if kljuc_masine(kljuc) == trazeni:
            return vrednost

    return default


def projekat_iz_masine(masina):
    projekat = pronadji_u_mapi(masina, MAPA_PROJEKATA)
    if projekat != "NEMAPIRANO":
        return projekat

    t = kljuc_masine(masina)

    if "APP350" in t:
        return "APP350"
    if "APP550" in t:
        return "APP550"
    if "LK4" in t or "LK3" in t:
        return "LK-4"
    if "EMR4" in t or "VITESCO" in t or "VITESKO" in t:
        return "VITESKO EMR4"
    if t.startswith("STATOR") or "APP550" in t:
        return "APP550"

    return "OSTALO"


def proces_iz_masine(masina):
    proces = pronadji_u_mapi(masina, MAPA_PROCESA)
    if proces != "NEMAPIRANO":
        return proces

    t = kljuc_masine(masina)

    if t.startswith("AIDA"):
        return "STAMPING"
    if t.startswith("ROTOR"):
        return "ROTOR"
    if t.startswith("STATOR") or t.startswith("HEATING"):
        return "WELDING"
    if t.startswith("DMC") and "GP12" in t:
        return "GP12"
    if t.startswith("DMC"):
        return "DMC"

    return "OSTALO"


def finalno_ocisti_razlog_od_imena(vrednost):
    tekst = str(vrednost or "")
    for rec in RECI_KOJE_SE_BRISU_IZ_RAZLOGA:
        tekst = re.sub(r"(?<!\w)" + re.escape(rec) + r"(?!\w)", " ", tekst, flags=re.IGNORECASE)
    tekst = re.sub(r"\s+", " ", tekst).strip()
    tekst = re.sub(r"^[\.\,\;\:\-\–\—\_\s]+", "", tekst).strip()
    tekst = re.sub(r"[\.\,\;\:\-\–\—\_\s]+$", "", tekst).strip()
    return re.sub(r"\s+", " ", tekst).strip()


def normalizuj_razlog(razlog):
    razlog = finalno_ocisti_razlog_od_imena(razlog)
    tekst = str(razlog or "").strip().lower()
    tekst = tekst.replace(":", " ").strip(" -–—:;,.`_")
    tekst = " ".join(tekst.split())
    if tekst == "":
        return "nije upisan razlog"
    tekst = re.sub(r"\botis(?:ak|am|ci)\b", "otisak", tekst, flags=re.IGNORECASE)
    return tekst


def ocisti_note(note):
    if note is None:
        return None
    tekst = str(note)
    for pojam in PUNA_IMENA_ZA_BRISANJE:
        tekst = re.sub(re.escape(pojam), "", tekst, flags=re.IGNORECASE)
    tekst = re.sub(r"\boperacije\b", "", tekst, flags=re.IGNORECASE)
    tekst = tekst.replace(":", " ").replace("\r", "\n").strip()
    linije = [linija.strip() for linija in tekst.split("\n") if linija.strip()]
    tekst = "\n".join(linije).strip()
    return tekst or None


def da_li_je_samo_ime_kolege(tekst):
    original = str(tekst or "")
    ocisceno = finalno_ocisti_razlog_od_imena(original)
    ocisceno = re.sub(r"^[\s\.\,\;\:\-\–\—\_`]+", "", ocisceno).strip()
    ocisceno = re.sub(r"[\s\.\,\;\:\-\–\—\_`]+$", "", ocisceno).strip()
    ima_ime = any(
        re.search(r"(?<!\w)" + re.escape(rec) + r"(?!\w)", original, flags=re.IGNORECASE)
        for rec in RECI_KOJE_SE_BRISU_IZ_RAZLOGA
    )
    return ima_ime and ocisceno == ""


def vrednost_merged_celije(ws, red, kolona):
    cell = ws.cell(row=red, column=kolona)
    if cell.value is not None:
        return cell.value
    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
    return None


def normalizuj_smenu(vrednost):
    tekst = ocisti_tekst(vrednost)
    if tekst in DOZVOLJENE_SMENE:
        return tekst
    return None


def smena_za_kolonu(ws, kolona):
    return normalizuj_smenu(vrednost_merged_celije(ws, RED_SMENA, kolona))


def poslednja_kolona_koja_pripada_smeni(ws):
    poslednja = None
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= RED_SMENA <= merged_range.max_row:
            smena = normalizuj_smenu(ws.cell(row=merged_range.min_row, column=merged_range.min_col).value)
            if smena is not None and (poslednja is None or merged_range.max_col > poslednja):
                poslednja = merged_range.max_col
    for kolona in range(1, min(ws.max_column, 160) + 1):
        smena = normalizuj_smenu(ws.cell(row=RED_SMENA, column=kolona).value)
        if smena is not None and (poslednja is None or kolona > poslednja):
            poslednja = kolona
    return poslednja or min(ws.max_column, 160)


def podeli_note_na_stavke(note):
    if note is None:
        return []
    tekst = str(note).replace("\r", "\n")
    delovi = [linija.strip() for linija in tekst.split("\n") if linija.strip()]
    if not delovi and tekst.strip():
        delovi = [tekst.strip()]
    return delovi


def parsiraj_stop_note(note, vrednost_celije):
    stavke = podeli_note_na_stavke(note)
    rezultat = []

    pattern_minuti = re.compile(
        r"(?<!\d)(\d+(?:[.,]\d+)?)\s*(?:minuta|minut|min|minute|mi\b|m\b|'|′)",
        re.IGNORECASE,
    )
    pattern_apostrof_pre = re.compile(r"(?:'|′)\s*(\d+(?:[.,]\d+)?)", re.IGNORECASE)
    pattern_broj_bez_jedinice = re.compile(
        r"^\s*[`'′\-–—]*\s*(\d+(?:[.,]\d+)?)\s+(.+\S)\s*$",
        re.IGNORECASE,
    )

    for stavka in stavke:
        stavka = str(stavka).strip()
        if not stavka:
            continue
        if da_li_je_samo_ime_kolege(stavka):
            continue

        m = pattern_minuti.search(stavka)
        koristi_apostrof_pre = False
        if not m:
            m = pattern_apostrof_pre.search(stavka)
            koristi_apostrof_pre = m is not None

        if m:
            minuti = broj(m.group(1).replace(",", "."))
            if koristi_apostrof_pre:
                razlog_original = pattern_apostrof_pre.sub("", stavka, count=1).strip(" -–—:;,.`")
            else:
                razlog_original = pattern_minuti.sub("", stavka, count=1).strip(" -–—:;,.`")
            razlog = normalizuj_razlog(razlog_original)
            if razlog != "nije upisan razlog":
                rezultat.append({"Trajanje_min": minuti, "Razlog": razlog, "Originalna_stavka": stavka})
            continue

        m = pattern_broj_bez_jedinice.match(stavka)
        if m and not re.search(r"\d{1,2}:\d{2}|\d{1,2}\.\d{2}", stavka):
            minuti = broj(m.group(1).replace(",", "."))
            razlog = normalizuj_razlog(m.group(2))
            if razlog != "nije upisan razlog":
                rezultat.append({"Trajanje_min": minuti, "Razlog": razlog, "Originalna_stavka": stavka})
            continue

        razlog = normalizuj_razlog(stavka)
        if razlog == "nije upisan razlog":
            continue
        rezultat.append({"Trajanje_min": broj(vrednost_celije), "Razlog": razlog, "Originalna_stavka": stavka})

    return rezultat


def jeste_plan_kolona(naziv):
    t = norm(naziv)
    return "plan" in t


def jeste_realizacija_kolona(naziv):
    t = norm(naziv)
    if any(x in t for x in ["nok", "scrap", "plan", "stop", "stops", "opening", "target"]):
        return False
    return any(x in t for x in [
        "realization", "realisation", "production", "good parts", "ok parts", "good parts/stator", "good parts/rotor"
    ])


def jeste_stop_kolona(naziv):
    t = norm(naziv)
    return "stop" in t and "min" in t


def datum_kao_date(vrednost):
    d = pd.to_datetime(vrednost, errors="coerce")
    if pd.isna(d):
        return None
    return d.date()

# ============================================================
# UČITAVANJE EXCEL-A
# ============================================================

@st.cache_data(show_spinner="Čitam Excel fajl i notes-e...")
def ucitaj_excel(upload_bytes):
    wb = load_workbook(BytesIO(upload_bytes), data_only=True)
    redovi = []
    zastoji = []

    for ws in wb.worksheets:
        masina = ws.title.strip()
        if masina in PRESKOCI_TABOVE:
            continue

        projekat = projekat_iz_masine(masina)
        proces = proces_iz_masine(masina)
        zadnja_kolona = poslednja_kolona_koja_pripada_smeni(ws)

        for red in range(PRVI_RED_PODATAKA, POSLEDNJI_RED_PODATAKA + 1):
            datum_raw = ws.cell(row=red, column=KOLONA_DATUM).value
            datum = datum_kao_date(datum_raw)
            if datum is None:
                continue

            po_smeni = {
                sm: {"Plan": 0.0, "Realizacija": 0.0, "Zastoj_min": 0.0}
                for sm in DOZVOLJENE_SMENE
            }

            for kolona in range(1, zadnja_kolona + 1):
                smena = smena_za_kolonu(ws, kolona)
                if smena is None:
                    continue

                naziv = ws.cell(row=RED_NAZIVI, column=kolona).value
                if naziv is None:
                    continue

                cell = ws.cell(row=red, column=kolona)
                vrednost = broj(cell.value)

                if jeste_plan_kolona(naziv):
                    po_smeni[smena]["Plan"] += vrednost

                if jeste_realizacija_kolona(naziv):
                    po_smeni[smena]["Realizacija"] += vrednost

                if jeste_stop_kolona(naziv):
                    po_smeni[smena]["Zastoj_min"] += vrednost

                    note = ocisti_note(cell.comment.text) if cell.comment else None
                    if note:
                        for stavka in parsiraj_stop_note(note, vrednost):
                            zastoji.append({
                                "Datum": datum,
                                "Masina": masina,
                                "Projekat": projekat,
                                "Proces": proces,
                                "Smena": smena,
                                "Smena_prikaz": SMENA_LABEL.get(smena, smena),
                                "Kolona": cell.coordinate,
                                "Razlog": stavka["Razlog"],
                                "Trajanje_min": stavka["Trajanje_min"],
                                "Originalna_stavka": stavka["Originalna_stavka"],
                            })

            for smena, vrednosti in po_smeni.items():
                plan = vrednosti["Plan"]
                realizacija = vrednosti["Realizacija"]
                redovi.append({
                    "Datum": datum,
                    "Masina": masina,
                    "Projekat": projekat,
                    "Proces": proces,
                    "Smena": smena,
                    "Smena_prikaz": SMENA_LABEL.get(smena, smena),
                    "Plan": plan,
                    "Realizacija": realizacija,
                    "Realizacija_%": procenat(realizacija, plan),
                    "Zastoj_min": vrednosti["Zastoj_min"],
                })

    df = pd.DataFrame(redovi)
    df_zastoji = pd.DataFrame(zastoji)

    if not df.empty:
        for kol in ["Plan", "Realizacija", "Zastoj_min"]:
            df[kol] = pd.to_numeric(df[kol], errors="coerce").fillna(0)

    if not df_zastoji.empty:
        df_zastoji["Trajanje_min"] = pd.to_numeric(df_zastoji["Trajanje_min"], errors="coerce").fillna(0)

    return df, df_zastoji


# ============================================================
# MODERNI HTML PRIKAZI
# ============================================================

def render_top_kpi(plan, realizacija, proc, zastoj):
    st.markdown(
        f"""
        <div class="neo-grid">
            <div class="neo-card"><div class="neo-label">Ukupan plan</div><div class="neo-value">{format_broj(plan)}</div></div>
            <div class="neo-card"><div class="neo-label">Ukupna realizacija</div><div class="neo-value">{format_broj(realizacija)}</div></div>
            <div class="neo-card"><div class="neo-label">Realizacija</div><div class="neo-value">{format_proc(proc)}</div></div>
            <div class="neo-card"><div class="neo-label">Ukupan zastoj</div><div class="neo-value">{format_broj(zastoj)} min</div></div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_summary_cards(summary):
    if summary is None or summary.empty:
        st.info("Nema zbirnih podataka za prikaz.")
        return

    cards = []
    for _, r in summary.sort_values("Masina").iterrows():
        cards.append(
            f"""
            <div class="summary-card">
                <div class="summary-title">🏭 {escape(str(r['Masina']))}</div>
                <div><span class="neo-pill">{escape(str(r['Projekat']))}</span><span class="neo-pill">{escape(str(r['Proces']))}</span></div>
                <div class="neo-sub">Plan: <b>{format_broj(r['Plan'])}</b> · Realizacija: <b>{format_broj(r['Realizacija'])}</b></div>
                <div class="neo-sub">Realizacija: <b>{format_proc(r['Realizacija_%'])}</b> · Zastoj: <b>{format_broj(r['Zastoj_min'])} min</b></div>
            </div>
            """
        )
    st.markdown('<div class="summary-grid">' + "".join(cards) + '</div>', unsafe_allow_html=True)


def render_smena_cards(df_tabela):
    cards = []
    for _, r in df_tabela.iterrows():
        total_cls = " neo-card-total" if str(r["Smena"]).upper() == "UKUPNO" else ""
        cards.append(
            f"""
            <div class="neo-card{total_cls}">
                <div class="neo-label">{escape(str(r['Smena']))}</div>
                <div class="neo-sub">Plan</div><div class="neo-value">{format_broj(r['Plan'])}</div>
                <div class="neo-sub">Realizacija: <b>{format_broj(r['Realizacija'])}</b> · <b>{format_proc(r['Realizacija %'])}</b></div>
                <div class="neo-sub">Zastoj: <b>{format_broj(r['Zastoj/min'])} min</b></div>
            </div>
            """
        )
    st.markdown('<div class="neo-grid">' + "".join(cards) + '</div>', unsafe_allow_html=True)


def render_stop_cards(df_stop, prikazi_originalne_stavke=False):
    if df_stop is None or df_stop.empty:
        st.info("Nema pročitanih razloga zastoja iz notes-a za ovu mašinu i datum.")
        return

    cards = []
    for _, r in df_stop.sort_values(["Smena_prikaz", "Kolona", "Razlog"]).iterrows():
        original = ""
        if prikazi_originalne_stavke:
            original = f"<div class='stop-meta'>Original: {escape(str(r.get('Originalna_stavka', '')))}</div>"
        cards.append(
            f"""
            <div class="stop-card">
                <div class="stop-reason">{escape(str(r['Razlog']))}</div>
                <div class="stop-meta"><b>{escape(str(r['Smena_prikaz']))}</b> · {format_broj(r['Trajanje_min'])} min · ćelija {escape(str(r['Kolona']))}</div>
                {original}
            </div>
            """
        )
    st.markdown('<div class="stop-grid">' + "".join(cards) + '</div>', unsafe_allow_html=True)

# ============================================================
# UI
# ============================================================

st.title("📊 Dnevni pregled proizvodnje po mašinama")
st.caption("Online verzija: korisnik sam učitava Excel fajl. Fajl se ne čuva u aplikaciji.")

with st.sidebar:
    st.header("📁 Učitavanje")
    uploaded_file = st.file_uploader(
        "Učitaj Production realization Excel fajl",
        type=["xlsx"],
        help="Izaberi mesečni Excel fajl sa firminog servera ili računara.",
    )

if uploaded_file is None:
    st.info("Učitaj Excel fajl da bi se prikazao dnevni pregled proizvodnje.")
    st.markdown(
        """
        **Kako se koristi:**
        1. Otvori link aplikacije.
        2. Sa leve strane klikni **Browse files**.
        3. Izaberi mesečni Excel fajl, na primer `06. Production realization JUN.xlsx`.
        4. Izaberi datum, mašinu, projekat i proces.
        """
    )
    st.stop()

try:
    upload_bytes = uploaded_file.getvalue()
    df, df_zastoji = ucitaj_excel(upload_bytes)
except Exception as e:
    st.error("Excel fajl nije mogao da se pročita.")
    st.exception(e)
    st.stop()

if df.empty:
    st.warning("Nisu pronađeni redovi za proizvodnju. Proveri da li je struktura fajla ista kao u postojećem programu.")
    st.stop()

svi_datumi = sorted(df["Datum"].dropna().unique())
min_datum = min(svi_datumi)
max_datum = max(svi_datumi)

with st.sidebar:
    st.header("🔎 Filteri")
    izabrani_datum = st.date_input("Datum", value=max_datum, min_value=min_datum, max_value=max_datum)

    projekti = ["Svi"] + sorted(df["Projekat"].dropna().unique().tolist())
    izabrani_projekat = st.selectbox("Projekat", projekti)

    procesi = ["Svi"] + sorted(df["Proces"].dropna().unique().tolist())
    izabrani_proces = st.selectbox("Proces", procesi)

    df_za_masine = df.copy()
    if izabrani_projekat != "Svi":
        df_za_masine = df_za_masine[df_za_masine["Projekat"] == izabrani_projekat]
    if izabrani_proces != "Svi":
        df_za_masine = df_za_masine[df_za_masine["Proces"] == izabrani_proces]

    masine = ["Sve"] + sorted(df_za_masine["Masina"].dropna().unique().tolist())
    izabrane_masine = st.multiselect("Mašina", masine, default=["Sve"])

    prikazi_originalne_stavke = st.checkbox("Prikaži originalne stavke iz notes-a", value=False)

maska = df["Datum"] == izabrani_datum
if izabrani_projekat != "Svi":
    maska &= df["Projekat"] == izabrani_projekat
if izabrani_proces != "Svi":
    maska &= df["Proces"] == izabrani_proces
if "Sve" not in izabrane_masine and izabrane_masine:
    maska &= df["Masina"].isin(izabrane_masine)

df_f = df[maska].copy()

maska_z = df_zastoji["Datum"] == izabrani_datum if not df_zastoji.empty else pd.Series(dtype=bool)
if not df_zastoji.empty:
    if izabrani_projekat != "Svi":
        maska_z &= df_zastoji["Projekat"] == izabrani_projekat
    if izabrani_proces != "Svi":
        maska_z &= df_zastoji["Proces"] == izabrani_proces
    if "Sve" not in izabrane_masine and izabrane_masine:
        maska_z &= df_zastoji["Masina"].isin(izabrane_masine)
    df_z_f = df_zastoji[maska_z].copy()
else:
    df_z_f = pd.DataFrame()

if df_f.empty:
    st.warning("Nema podataka za izabrane filtere.")
    st.stop()

uk_plan = df_f["Plan"].sum()
uk_real = df_f["Realizacija"].sum()
uk_zastoj = df_f["Zastoj_min"].sum()
uk_proc = procenat(uk_real, uk_plan)

render_top_kpi(uk_plan, uk_real, uk_proc, uk_zastoj)

st.divider()

# Kratak zbirni prikaz po mašinama
summary = (
    df_f.groupby(["Masina", "Projekat", "Proces"], as_index=False)
    .agg(Plan=("Plan", "sum"), Realizacija=("Realizacija", "sum"), Zastoj_min=("Zastoj_min", "sum"))
)
summary["Realizacija_%"] = summary.apply(lambda r: procenat(r["Realizacija"], r["Plan"]), axis=1)

st.subheader("📌 Pregled po mašinama")
render_summary_cards(summary[["Masina", "Projekat", "Proces", "Plan", "Realizacija", "Realizacija_%", "Zastoj_min"]])

# Detalji po mašinama
for masina in sorted(df_f["Masina"].unique()):
    df_m = df_f[df_f["Masina"] == masina].copy()
    if df_m.empty:
        continue

    projekat = df_m["Projekat"].iloc[0]
    proces = df_m["Proces"].iloc[0]

    st.markdown('<div class="machine-card">', unsafe_allow_html=True)
    st.markdown(f'<div class="machine-title">🏭 {masina}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="machine-subtitle">Projekat: <b>{projekat}</b> · Proces: <b>{proces}</b> · Datum: <b>{izabrani_datum.strftime("%d.%m.%Y")}</b></div>', unsafe_allow_html=True)

    redovi_tabele = []
    for smena in DOZVOLJENE_SMENE:
        r = df_m[df_m["Smena"] == smena]
        if r.empty:
            plan = real = zastoj = 0.0
        else:
            plan = r["Plan"].sum()
            real = r["Realizacija"].sum()
            zastoj = r["Zastoj_min"].sum()
        redovi_tabele.append({
            "Smena": SMENA_LABEL.get(smena, smena),
            "Plan": plan,
            "Realizacija": real,
            "Realizacija %": procenat(real, plan),
            "Zastoj/min": zastoj,
        })

    ukupno_plan = sum(x["Plan"] for x in redovi_tabele)
    ukupno_real = sum(x["Realizacija"] for x in redovi_tabele)
    ukupno_zastoj = sum(x["Zastoj/min"] for x in redovi_tabele)
    redovi_tabele.append({
        "Smena": "UKUPNO",
        "Plan": ukupno_plan,
        "Realizacija": ukupno_real,
        "Realizacija %": procenat(ukupno_real, ukupno_plan),
        "Zastoj/min": ukupno_zastoj,
    })

    df_tabela = pd.DataFrame(redovi_tabele)
    render_smena_cards(df_tabela)

    st.markdown('<div class="mini-title">⏱️ Razlozi zastoja iz notes-a</div>', unsafe_allow_html=True)
    df_z_m = df_z_f[df_z_f["Masina"] == masina].copy() if not df_z_f.empty else pd.DataFrame()

    render_stop_cards(df_z_m, prikazi_originalne_stavke=prikazi_originalne_stavke)

    st.markdown('</div>', unsafe_allow_html=True)
